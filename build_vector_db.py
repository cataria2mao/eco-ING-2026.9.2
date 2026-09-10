import os, json
from pathlib import Path
import chromadb
from docx import Document
from typing import List
from langchain_text_splitters import RecursiveCharacterTextSplitter
from dotenv import load_dotenv

load_dotenv(override=True)

SOP_DOCX_PATH = os.getenv("SOP_DOCX_PATH", r"D:\PythonProject1\生态环境调查报告工作.docx")
CHROMA_DIR = os.getenv("CHROMA_DIR", str(Path(__file__).resolve().parent / "chroma_sop_db"))
EMBED_MODEL = os.getenv("EMBED_MODEL", "text-embedding-v3")
EMBED_BASE_URL = os.getenv("EMBED_BASE_URL", "https://dashscope.aliyuncs.com/compatible-mode/v1")
EMBED_API_KEY = os.getenv("DASHSCOPE_API_KEY", "")
COLLECTION_NAME = "sop_semantic"


class DashScopeEmbeddingFunction:
    def __init__(self, model=EMBED_MODEL, api_key="", base_url=EMBED_BASE_URL):
        self.model = model
        self.api_key = api_key or EMBED_API_KEY
        self.base_url = base_url

    def _embed_batch(self, texts):
        import requests
        out = []
        for i in range(0, len(texts), 10):
            batch = texts[i:i + 10]
            r = requests.post(
                f"{self.base_url}/embeddings",
                headers={
                    "Authorization": f"Bearer {self.api_key}",
                    "Content-Type": "application/json",
                },
                json={"model": self.model, "input": batch},
                timeout=60,
            )
            r.raise_for_status()
            data = sorted(r.json().get("data", []), key=lambda x: x.get("index", 0))
            out.extend(d["embedding"] for d in data)
        return out

    def __call__(self, input):
        if isinstance(input, str):
            input = [input]
        return self._embed_batch(list(input))

    def embed_query(self, input):
        return self.__call__(input)

    def name(self):
        return "dashscope_text_embedding_v3"


def read_sop_chunks() -> List[str]:
    """读取 SOP docx，按段落+表格抽取文本，再切块。"""
    if not os.path.exists(SOP_DOCX_PATH):
        print(f"[VEC] SOP 文件不存在: {SOP_DOCX_PATH}")
        return []

    try:
        doc = Document(SOP_DOCX_PATH)
    except Exception as e:
        print(f"[VEC] 打开 SOP docx 失败: {e}")
        return []

    lines: List[str] = []

    # 1) 普通段落
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            lines.append(t)

    # 2) 表格（SOP 里常有关键表格，别漏掉）
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    full_text = "\n".join(lines)
    if not full_text.strip():
        print("[VEC] SOP 文档解析后为空，请检查文件内容")
        return []

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=50,
        separators=["\n\n", "\n", "。", "；", "，", " ", ""],
        length_function=len,   # 不依赖 tiktoken
    )
    chunks = splitter.split_text(full_text)
    print(f"[VEC] SOP 读取完成：{len(lines)} 段 → {len(chunks)} 个文本块")
    return chunks


def main():
    chunks = read_sop_chunks()
    if not chunks:
        print("没有读取到 SOP 分块，终止。")
        return

    Path(CHROMA_DIR).mkdir(parents=True, exist_ok=True)

    client = chromadb.PersistentClient(path=CHROMA_DIR)

    # 只有重新构建时才删除旧 collection
    try:
        client.delete_collection(COLLECTION_NAME)
    except Exception:
        pass

    ef = DashScopeEmbeddingFunction()
    coll = client.create_collection(
        name=COLLECTION_NAME,
        embedding_function=ef,
    )

    vectors = ef(chunks)
    coll.add(
        ids=[f"c{i}" for i in range(len(chunks))],
        documents=chunks,
        embeddings=vectors,
    )

    # 可选但推荐：保存 chunks，供原文件构建 BM25 用
    cache_file = Path(CHROMA_DIR) / "sop_chunks.json"
    cache_file.write_text(json.dumps(chunks, ensure_ascii=False), encoding="utf-8")

    print(f"构建完成：{len(chunks)} 个文本块 -> {CHROMA_DIR}")


if __name__ == "__main__":
    main()